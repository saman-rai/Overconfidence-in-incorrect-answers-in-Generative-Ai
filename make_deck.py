"""STEP 4 - Build the Week-2 progress deck from YOUR results.

Reads  results/summary.json + results/charts/*.png
Writes Group7_Week2_Progress_SamanRai.pptx  (every number comes from your run)

  python make_deck.py
  python make_deck.py --date "16 October 2026" --out my_deck.pptx
"""
import argparse
import json
import os
from collections import Counter

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Presenter / project details
# ---------------------------------------------------------------------------
PRESENTER = "Saman Rai"
STUDENT_ID = "2317130"
GROUP = "Group 7"
SUPERVISOR = "Prof. Simeon"
TEAM = ("Ramauli Madan Bahadur • Tamang Sajan • Jaiswal Anand Bihari • Sunar Binod • Kaliraj Bikram • "
        "Rai Saman • Dahal Sushil • Subedi Chetana • Saad")
FOOTER = "Overconfidence in GenAI Answers  •  Seminar on AI"

# ---------------------------------------------------------------------------
# Theme (matches the Week-1 deck)
# ---------------------------------------------------------------------------
BG = RGBColor(0x0E, 0x18, 0x28)
CARD = RGBColor(0x1B, 0x2B, 0x43)
CARD2 = RGBColor(0x16, 0x24, 0x3A)
BORDER = RGBColor(0x2F, 0x41, 0x5C)
INK = RGBColor(0xF5, 0xF8, 0xFC)
MUTED = RGBColor(0xB0, 0xBE, 0xD2)
CYAN = RGBColor(0x2D, 0xC4, 0xD3)
GREEN = RGBColor(0x5C, 0xC9, 0x8B)
YELLOW = RGBColor(0xF2, 0xC2, 0x48)
RED = RGBColor(0xEB, 0x68, 0x68)
FONT = "Aptos"

W, H = 13.333, 7.5
SLIDE_COUNT = [0]


def rgb(hexstr):
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def pct(x, d=0):
    return "n/a" if x is None else f"{x * 100:.{d}f}%"


def f3(x):
    return "n/a" if x is None else f"{x:.3f}"


def f2(x):
    return "n/a" if x is None else f"{x:.2f}"


def trunc(s, n):
    s = " ".join(str(s).split())
    return s if len(s) <= n else s[: n - 1].rstrip() + "…"


# ---------------------------------------------------------------------------
# Drawing helpers
# ---------------------------------------------------------------------------

def box(slide, x, y, w, h, fill=CARD, line=BORDER, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def text(slide, x, y, w, h, content, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0, space_after=0, italic=False):
    """content: str | list of paragraphs; a paragraph is str or list of (text, {opts}) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for rtext, opts in runs:
            r = p.add_run()
            r.text = rtext
            f = r.font
            f.name = FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = opts.get("color", color)
    return tb


def label(slide, x, y, w, s, color=CYAN, size=11):
    return text(slide, x, y, w, 0.3, s.upper(), size=size, color=color, bold=True)


def new_slide(prs, title, subtitle=None, notes=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    SLIDE_COUNT[0] += 1
    n = SLIDE_COUNT[0]
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if title:
        text(slide, 0.6, 0.38, 11.2, 0.6, title, size=28, bold=True)
        if subtitle:
            text(slide, 0.62, 0.98, 11.2, 0.35, subtitle, size=13, color=MUTED)
        text(slide, 12.2, 0.42, 0.55, 0.3, str(n), size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, 0.6, 7.03, 8, 0.25, FOOTER, size=9, color=MUTED)
    text(slide, 9.7, 7.03, 3.05, 0.25, DATE_SHORT[0], size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def chip(slide, x, y, w, s, color):
    b = box(slide, x, y, w, 0.34, fill=CARD2, line=color, radius=0.17)
    text(slide, x, y, w, 0.34, s, size=10.5, color=color, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return b


def picture_fit(slide, path, x, y, max_w, max_h):
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w, h = iw * scale, ih * scale
    return slide.shapes.add_picture(path, Inches(x + (max_w - w) / 2), Inches(y + (max_h - h) / 2), Inches(w), Inches(h))


DATE_SHORT = [""]
PROV = {"openai": "OpenAI", "gemini": "Google Gemini", "groq": "Groq", "openrouter": "OpenRouter"}

# ---------------------------------------------------------------------------
# Data-driven sentences (only facts computed from summary.json)
# ---------------------------------------------------------------------------

def findings(S):
    M = S["methods"]
    names = list(M)
    out = []
    v = M["Verbalized confidence"]["metrics"]
    vts = M["Verbalized + TS"]["metrics"]
    out.append(
        f"The model answered {pct(v['accuracy'])} correctly but claimed {pct(v['mean_conf'])} confidence on average "
        f"(gap {v['overconfidence_gap'] * 100:+.0f} pts)."
        if v["overconfidence_gap"] > 0.02 else
        f"Average verbalized confidence ({pct(v['mean_conf'])}) was close to accuracy ({pct(v['accuracy'])}) overall — "
        f"but the problem shows up in the wrong answers."
    )
    if v["wrong_with_high_conf_share"] is not None:
        out.append(f"{pct(v['wrong_with_high_conf_share'])} of the WRONG answers were stated with ≥90% confidence "
                   f"(avg. {pct(v['mean_conf_when_wrong'])} when wrong vs {pct(v['mean_conf_when_right'])} when right).")
    if vts["ece"] < v["ece"]:
        out.append(f"Temperature scaling reduced verbalized ECE from {f3(v['ece'])} to {f3(vts['ece'])} "
                   f"({(1 - vts['ece'] / v['ece']) * 100:.0f}% lower) — calibration can be repaired cheaply.")
    else:
        out.append(f"Temperature scaling did not lower verbalized ECE ({f3(v['ece'])} → {f3(vts['ece'])}).")
    au = {n: M[n]["metrics"]["auroc"] for n in names if M[n]["metrics"]["auroc"] is not None and not n.endswith("+ TS")}
    if au:
        best = max(au, key=au.get)
        out.append(f"Best at telling right from wrong answers: {best} (AUROC {f2(au[best])}). "
                   "Temperature scaling barely moves AUROC — it rescales scores, it does not re-rank answers.")
    return out


def verbal_rounding(S_raw_counts):
    return S_raw_counts


# ---------------------------------------------------------------------------
# Slides
# ---------------------------------------------------------------------------

def build(S, res_dir, out_path, date_long):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    run = S["run"]
    M = S["methods"]
    names = list(M)
    sc_name = [n for n in names if n.startswith("Self-consistency")][0]
    k = run.get("k", 5)
    N = run["n_questions"]
    model = run.get("model", "LLM")
    v = M["Verbalized confidence"]["metrics"]
    vts = M["Verbalized + TS"]["metrics"]
    sc = M[sc_name]["metrics"]
    tri = S["triage_main"]
    agree_needed = round(tri["tau"] * k)
    charts = os.path.join(res_dir, "charts")

    # ---------------- 1. Title
    s = new_slide(prs, None, notes=(
        f"Good morning. I'm {PRESENTER}, presenting Group 7's Week-2 progress on overconfidence in GenAI answers. "
        "Last week we defined the problem, reviewed the literature and proposed a tiered confidence-triage pipeline. "
        "This week we moved from planning to evidence: we built the evaluation harness, implemented the baselines we "
        "promised, and ran our first experiment on a real benchmark."))
    text(s, 0.75, 0.85, 10, 0.35, "WEEKLY PROGRESS REPORT  •  WEEK 2", size=13, color=CYAN, bold=True)
    text(s, 0.75, 1.35, 11.5, 0.9, "Overconfidence in GenAI Answers", size=40, bold=True)
    text(s, 0.77, 2.3, 11.5, 0.5, "First experiments: how often is an LLM confidently wrong?", size=18, color=MUTED)
    cards = [("Presenter", f"{PRESENTER}  •  ID {STUDENT_ID}"), ("Group & supervisor", f"{GROUP}  •  {SUPERVISOR}"),
             ("Current milestone", "Week 2 • Baselines + first results")]
    cw = 3.78
    for i, (lab, val) in enumerate(cards):
        x = 0.75 + i * (cw + 0.2)
        box(s, x, 3.3, cw, 1.25)
        label(s, x + 0.28, 3.52, cw - 0.5, lab, color=CYAN if i != 2 else GREEN)
        text(s, x + 0.28, 3.92, cw - 0.5, 0.45, val, size=15)
    text(s, 0.75, 4.9, 11.8, 0.5, "Team: " + TEAM, size=11, color=MUTED)
    text(s, 0.75, 5.75, 6, 0.4, date_long, size=13, bold=True)

    # ---------------- 2. Where we left off
    s = new_slide(prs, "1. Where We Left Off", "Week-1 commitment → what we delivered in Week 2", notes=(
        "Quick recap. Our research question is whether a tiered confidence-triage pipeline can reduce overconfident wrong "
        "answers at a practical cost. At the end of Week 1 we committed to implementing temperature scaling and verbalized "
        "confidence and getting first results. All of that is done. We also added self-consistency, because it is "
        f"Stage 1 of our own pipeline, so we can already get an early signal on whether triage is promising. "
        f"The experiment covered {N} questions."))
    box(s, 0.6, 1.55, 5.6, 5.15)
    label(s, 0.9, 1.8, 5, "Last week we committed to")
    text(s, 0.9, 2.2, 5.0, 1.6, "“Implement the baseline calibration methods — temperature scaling and verbalized "
         "confidence — and establish the first evaluation results.”", size=17, italic=True)
    label(s, 0.9, 4.05, 5, "Research question", color=YELLOW)
    text(s, 0.9, 4.45, 5.0, 1.6, "Can a tiered confidence-triage pipeline reduce overconfident wrong answers while "
         "keeping computation practical?", size=15, color=MUTED)
    rows = [("Dataset selected & sampled", f"DONE • {N} Qs", GREEN),
            ("Evaluation harness (grading + metrics)", "DONE", GREEN),
            ("Verbalized confidence baseline", "DONE", GREEN),
            ("Temperature scaling baseline", "DONE", GREEN),
            ("Self-consistency (Stage 1 of triage)", "DONE • bonus", GREEN),
            ("First evaluation results", "DONE", GREEN)]
    for i, (r, st, c) in enumerate(rows):
        y = 1.55 + i * 0.87
        box(s, 6.45, y, 4.35, 0.7)
        text(s, 6.7, y, 4.0, 0.7, r, size=14, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        box(s, 10.95, y, 1.8, 0.7)
        text(s, 10.95, y, 1.8, 0.7, st, size=12, color=c, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # ---------------- 3. What we built
    s = new_slide(prs, "2. What We Built This Week", "An end-to-end, reproducible evaluation harness", notes=(
        "This is the harness we built. Step one samples questions from the benchmark with a fixed random seed so results are "
        "reproducible. Step two queries the model three ways for every question. Step three grades answers automatically "
        "against all accepted gold answers. Step four fits temperature scaling with 5-fold cross-validation, so we never "
        "evaluate on the data we calibrated on. Step five computes all the metrics from our Week-1 plan. The key point: "
        "every method sees the same questions, grader and metrics, so in Week 3 our triage pipeline just plugs in as one more method."))
    steps = [("1", "Sample", f"{N} questions, fixed random seed"),
             ("2", "Query LLM", f"Greedy answer,\n{k} samples,\nverbalized conf."),
             ("3", "Auto-grade", "Normalized match against all gold answer aliases"),
             ("4", "Calibrate", "Temperature scaling,\n5-fold cross-fitted"),
             ("5", "Measure", "Accuracy, ECE, Brier, AUROC, risk–coverage")]
    bw, gap = 2.15, 0.36
    for i, (num, head, body) in enumerate(steps):
        x = 0.6 + i * (bw + gap)
        box(s, x, 1.65, bw, 2.55)
        b = box(s, x + 0.25, 1.9, 0.5, 0.5, fill=BG, line=CYAN, radius=0.25)
        text(s, x + 0.25, 1.9, 0.5, 0.5, num, size=15, color=CYAN, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, 2.6, bw - 0.45, 0.4, head, size=17, bold=True)
        text(s, x + 0.25, 3.08, bw - 0.4, 1.05, body, size=12.5, color=MUTED)
        if i < len(steps) - 1:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.06), Inches(2.8), Inches(0.24), Inches(0.26))
            a.fill.solid()
            a.fill.fore_color.rgb = CYAN
            a.line.fill.background()
    box(s, 0.6, 4.55, 5.95, 2.15)
    label(s, 0.9, 4.78, 5.4, "Why this matters", color=GREEN)
    text(s, 0.9, 5.18, 5.4, 1.4, ["Same questions, grader and metrics for every method → fair comparison.",
                                  "Week 3: our triage pipeline plugs in as just one more method."], size=13.5, space_after=6)
    box(s, 6.8, 4.55, 5.95, 2.15)
    label(s, 7.1, 4.78, 5.4, "Reproducibility", color=YELLOW)
    text(s, 7.1, 5.18, 5.4, 1.4, ["Every raw model output is logged (one JSON line per question).",
                                  "Runs are resumable and rate-limit safe; one command re-creates every number and chart."],
         size=13.5, space_after=6)

    # ---------------- 4. Setup
    mins = run.get("wall_seconds", 0) / 60
    tokens = run.get("prompt_tokens", 0) + run.get("completion_tokens", 0)
    s = new_slide(prs, "3. Experimental Setup", "What exactly we ran", notes=(
        f"Setup details. We used {run.get('dataset_label')} — factual questions with ground-truth answers and several accepted "
        f"aliases, which makes automatic grading reliable. We sampled {N} questions with a fixed seed. The model is {model}, "
        f"accessed through the {run.get('provider', '')} API. For each question we made a greedy answer at temperature 0"
        f"{' with token log-probabilities' if run.get('logprobs_available') else ''}, {k} sampled answers at temperature 1, and a "
        "separate prompt asking for an answer plus a confidence percentage. "
        f"In total {run.get('api_calls', 0):,} API calls and about {tokens:,} tokens — cheap enough to scale up next week."))
    cells = [("Dataset", run.get("dataset_label", ""), f"{N} questions • random sample (seed {run.get('seed', 7)})"),
             ("Model", trunc(model, 32), f"via {PROV.get(run.get('provider', ''), run.get('provider', ''))} API • token log-probs: "
                                          f"{'yes' if run.get('logprobs_available') else 'not available'}"),
             ("Probes per question", f"1 greedy + {k} samples + 1 verbalized",
              "Greedy & verbalized at temperature 0; samples at temperature 1.0"),
             ("Grading", "Automatic, lenient alias match",
              "Lower-case, strip punctuation/articles; correct if a gold alias appears in the answer"),
             ("Calibration", "Temperature scaling (1 parameter)", "Fitted on 4/5 of questions, applied to the held-out 1/5 (×5 folds)"),
             ("Run cost", f"{run.get('api_calls', 0):,} API calls", f"≈ {tokens:,} tokens • {('<1' if mins < 1 else f'{mins:.0f}')} min wall-clock")]
    cw, ch = 3.9, 2.45
    for i, (lab, big, small) in enumerate(cells):
        x = 0.6 + (i % 3) * (cw + 0.25)
        y = 1.55 + (i // 3) * (ch + 0.25)
        box(s, x, y, cw, ch)
        label(s, x + 0.3, y + 0.28, cw - 0.55, lab, color=[CYAN, CYAN, YELLOW, YELLOW, GREEN, GREEN][i])
        text(s, x + 0.3, y + 0.7, cw - 0.55, 0.8, big, size=19, bold=True)
        text(s, x + 0.3, y + 1.5, cw - 0.55, 0.85, small, size=13.5, color=MUTED)

    # ---------------- 5. Methods
    s = new_slide(prs, "4. Confidence Signals Compared", "Every method attaches a confidence score in [0, 1] to an answer",
                  notes=("These are the confidence signals we compared. "
                         + ("Token probability is the probability the model itself assigns to its answer tokens. " if "Token probability" in M else "")
                         + "Verbalized confidence simply asks the model how sure it is. Temperature scaling learns a single "
                         "number T from held-out data and rescales the confidence; when T is above one it pulls overconfident scores "
                         f"down. Self-consistency asks the same question {k} times with randomness and checks how often the answers agree. "
                         "Note the cost column: self-consistency is the most expensive, which is exactly why our pipeline uses it only as a "
                         "cheap first filter and saves heavy verification for the uncertain cases."))
    hdr_y = 1.6
    cols = [(0.6, 3.3, "Method"), (4.0, 6.6, "How the confidence is obtained"), (10.75, 2.0, "Generations / Q")]
    for x, w, h in cols:
        label(s, x + 0.2, hdr_y, w, h, color=MUTED, size=10.5)
    row_h = min(0.88, (4.3 - 0.12 * (len(names) - 1)) / len(names))
    for i, n in enumerate(names):
        y = 2.0 + i * (row_h + 0.12)
        box(s, 0.6, y, 12.15, row_h)
        swatch = box(s, 0.6, y, 0.09, row_h, fill=rgb(M[n]["color"]), line=None, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.85, y, 3.1, row_h, n, size=14.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.2, y, 6.4, row_h, M[n]["description"], size=12.5, color=MUTED, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 10.75, y, 2.0, row_h, str(M[n]["api_calls_per_question"]), size=15, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.6, 6.52, 12.1, 0.35, [[("Temperature scaling: ", {"bold": True, "color": CYAN}),
                                     ("calibrated = σ( logit(p) / T ).  T > 1 softens over-confident scores; T is learned from held-out data.", {})]],
         size=12.5)

    # ---------------- 6. Headline results
    s = new_slide(prs, "5. Headline Results", f"{N} questions • {trunc(model, 30)} • {run.get('dataset_label', '').split(' (')[0]}", notes=(
        "Here are the headline numbers. Accuracy is how often the answer is right; average confidence is what the model claims. "
        "ECE — expected calibration error — is the gap between claimed confidence and actual accuracy; lower is better. Brier score "
        "also measures probability quality; lower is better. AUROC measures whether confidence separates right from wrong answers; "
        "0.5 is random, 1.0 is perfect. The last column is our headline overconfidence measure: of the answers that were WRONG, what "
        f"share were stated with at least 90% confidence. For verbalized confidence that is {pct(v['wrong_with_high_conf_share'])}. "
        "Green marks the best value in each column."))
    heads = ["Method", "Accuracy", "Avg. conf.", "ECE ↓", "Brier ↓", "AUROC ↑", "Wrong but ≥90% conf. ↓"]
    keys = [None, "accuracy", "mean_conf", "ece", "brier", "auroc", "wrong_with_high_conf_share"]
    xs = [0.6, 3.85, 5.15, 6.45, 7.75, 9.05, 10.35]
    ws = [3.25, 1.3, 1.3, 1.3, 1.3, 1.3, 2.4]
    best = {}
    for kk, better in [("ece", min), ("brier", min), ("auroc", max), ("wrong_with_high_conf_share", min)]:
        vals = [M[n]["metrics"][kk] for n in names if M[n]["metrics"][kk] is not None]
        if vals:
            best[kk] = better(vals)
    for x, w, h in zip(xs, ws, heads):
        text(s, x + (0.2 if h == "Method" else 0), 1.55, w, 0.3, h, size=11, color=MUTED, bold=True,
             align=PP_ALIGN.LEFT if h == "Method" else PP_ALIGN.CENTER)
    rh = 0.5
    for i, n in enumerate(names):
        y = 1.95 + i * (rh + 0.08)
        box(s, 0.6, y, 12.15, rh, fill=CARD if i % 2 == 0 else CARD2)
        box(s, 0.6, y, 0.09, rh, fill=rgb(M[n]["color"]), line=None, radius=0, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.85, y, 3.0, rh, n, size=13.5, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        mm = M[n]["metrics"]
        for kk, x, w in zip(keys[1:], xs[1:], ws[1:]):
            val = mm[kk]
            if kk in ("accuracy", "mean_conf", "wrong_with_high_conf_share"):
                sv = pct(val, 1)
            elif kk == "auroc":
                sv = f2(val)
            else:
                sv = f3(val)
            col = GREEN if (kk in best and val is not None and abs(val - best[kk]) < 1e-9) else INK
            text(s, x, y, w, rh, sv, size=14, color=col, bold=(col == GREEN), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    ty = 1.95 + len(names) * (rh + 0.08) + 0.3
    kh = min(1.9, max(1.2, 6.75 - ty))
    kpis = [(pct(v["wrong_with_high_conf_share"]), "of wrong answers came with ≥90% verbalized confidence", RED),
            (f"{f3(v['ece'])} → {f3(vts['ece'])}", "verbalized ECE before → after temperature scaling",
             GREEN if vts["ece"] < v["ece"] else YELLOW),
            (f2(sc["auroc"]), f"AUROC of self-consistency ({k} samples) at flagging wrong answers", CYAN)]
    kw = 3.9
    for i, (big, small, c) in enumerate(kpis):
        x = 0.6 + i * (kw + 0.225)
        box(s, x, ty, kw, kh)
        text(s, x + 0.3, ty + 0.18, kw - 0.6, 0.6, big, size=26, bold=True, color=c)
        text(s, x + 0.3, ty + 0.8, kw - 0.6, kh - 0.9, small, size=12, color=MUTED)

    # ---------------- 7. Reliability diagrams
    worst = max((n for n in names), key=lambda n: M[n]["metrics"]["overconfidence_gap"])
    s = new_slide(prs, "6. Calibration: Stated Confidence vs. Reality", "Reliability diagrams — a perfectly calibrated model sits on the dashed diagonal",
                  notes=("How to read these: we group answers by stated confidence, then plot how often each group was actually correct. "
                         "On the dashed diagonal, 80% confidence means 80% accuracy. Bars BELOW the diagonal mean the model was less accurate "
                         "than it claimed — that is overconfidence. Look at where most answers sit and how far they are from the line, and "
                         "compare the raw verbalized panel with the temperature-scaled panel. "
                         f"In our run the largest average overconfidence gap was {worst}, at {M[worst]['metrics']['overconfidence_gap'] * 100:+.0f} percentage points."))
    picture_fit(s, os.path.join(charts, "reliability.png"), 0.6, 1.5, 12.15, 4.3)
    box(s, 0.6, 5.95, 12.15, 0.8)
    text(s, 0.9, 5.95, 11.6, 0.8, [[("Reading it: ", {"bold": True, "color": CYAN}),
                                    ("bars below the diagonal = over-confident. ", {}),
                                    (f"Largest average gap: {worst} ({M[worst]['metrics']['overconfidence_gap'] * 100:+.0f} pts "
                                     f"confidence above accuracy).", {"color": MUTED})]], size=13.5, anchor=MSO_ANCHOR.MIDDLE)

    # ---------------- 8. Evidence
    ex = S["confidently_wrong_examples"][:3]
    s = new_slide(prs, "7. Evidence: Confidently Wrong Answers", "Distribution of verbalized confidence for right vs. wrong answers",
                  notes=("This chart is the core of our topic. Green bars are correct answers, red bars are wrong answers, by the confidence "
                         "the model stated. If the model knew what it didn't know, the red bars would sit on the left. Instead, "
                         f"{pct(v['wrong_with_high_conf_share'])} of the wrong answers were stated with at least 90% confidence. "
                         "On the right are real examples from our run — the model is wrong and still very sure. "
                         f"Overall we found {S['n_confidently_wrong_verbal']} confidently-wrong answers out of {N} questions."))
    picture_fit(s, os.path.join(charts, "verbal_conf_hist.png"), 0.5, 1.5, 6.6, 4.4)
    text(s, 0.6, 6.05, 6.4, 0.75, [[(f"{S['n_confidently_wrong_verbal']} of {N}", {"bold": True, "color": RED}),
                                    (" answers were wrong while the model claimed ≥90% confidence.", {})]], size=14)
    if ex:
        eh = min(1.65, (5.3 - 0.15 * (len(ex) - 1)) / len(ex))
        for i, e in enumerate(ex):
            y = 1.5 + i * (eh + 0.15)
            box(s, 7.35, y, 5.4, eh)
            text(s, 7.6, y + 0.15, 4.95, 0.62, trunc(e["question"], 105), size=11.5, color=MUTED)
            text(s, 7.6, y + eh - 0.72, 4.95, 0.3, [[("Model: ", {"color": MUTED}),
                                                      (trunc(e["model_answer"], 40), {"bold": True, "color": RED}),
                                                      (f"   ({pct(e['confidence'])} sure)", {"color": MUTED})]], size=12.5)
            text(s, 7.6, y + eh - 0.4, 4.95, 0.3, [[("Correct: ", {"color": MUTED}),
                                                     (trunc(e["gold"], 45), {"bold": True, "color": GREEN})]], size=12.5)
    else:
        box(s, 7.35, 1.5, 5.4, 2.0)
        text(s, 7.6, 1.7, 4.9, 1.6, "No answer in this run was wrong with ≥90% stated confidence.", size=14)

    # ---------------- 9. Triage preview
    s = new_slide(prs, "8. Early Look: Does Triage Make Sense?", f"Stage 1 of our pipeline = self-consistency filter ({k} samples)",
                  notes=(f"This is an early test of our own idea, using data we already have. Stage 1 of the pipeline accepts an answer if at least "
                         f"{agree_needed} of the {k} samples agree; otherwise it escalates to an expensive check. The curve shows error rate "
                         "against coverage — if we answer only the most confident questions, how many mistakes remain. A lower curve is better. "
                         f"At our threshold we would answer {pct(tri['accepted_share'])} of questions directly with "
                         f"{pct(tri['accuracy_accepted'])} accuracy, and escalate only {pct(tri['escalated_share'])}, which contain "
                         f"{pct(tri['errors_caught_share'])} of all errors. "
                         + ("So expensive verification would only be needed for a minority of queries. "
                            if tri["escalated_share"] < 0.5 else
                            "At this threshold a majority of queries would still be escalated, so tuning the threshold is part of next week's work. ")
                         + "Building and costing that Stage-2 check is our Week 3–4 work."))
    picture_fit(s, os.path.join(charts, "risk_coverage.png"), 0.5, 1.5, 6.6, 4.4)
    text(s, 0.6, 6.05, 6.4, 0.75, "Lower curve = fewer errors among the answers we choose to give. "
         "White dot = our triage threshold.", size=12, color=MUTED)
    stats = [(pct(tri["accepted_share"]), f"answered directly (≥{agree_needed}/{k} samples agree)",
              f"accuracy on these: {pct(tri['accuracy_accepted'])}", GREEN),
             (pct(tri["escalated_share"]), "escalated to Stage 2 (low agreement)",
              f"error rate on these: {pct(tri['error_rate_escalated'])}", YELLOW),
             (pct(tri["errors_caught_share"]), "of ALL wrong answers land in the escalated group",
              "→ verification effort goes where the errors are", CYAN)]
    for i, (big, mid, small, c) in enumerate(stats):
        y = 1.5 + i * 1.82
        box(s, 7.35, y, 5.4, 1.65)
        text(s, 7.65, y + 0.2, 1.9, 0.7, big, size=30, bold=True, color=c)
        text(s, 9.55, y + 0.25, 3.0, 0.7, mid, size=13, bold=True)
        text(s, 9.55, y + 1.0, 3.0, 0.5, small, size=11.5, color=MUTED)

    # ---------------- 10. Findings + lessons
    F = findings(S)
    s = new_slide(prs, "9. Key Findings & Lessons Learned", "What the first results tell us", notes=(
        "To summarise the findings: " + " ".join(F) + " Lessons learned: automatic grading of free-text answers is the biggest "
        "source of measurement error, so we spot-check it manually; and temperature scaling fixes the scale of confidence but not "
        "which answers get high confidence — that is why we need a stronger signal such as self-consistency plus verification, "
        "which is exactly our triage design."))
    box(s, 0.6, 1.55, 7.0, 5.15)
    label(s, 0.9, 1.78, 6.5, "Findings (from our data)", color=GREEN)
    text(s, 0.9, 2.25, 6.45, 4.35, [[("•  ", {"color": GREEN, "bold": True}), (f_, {})] for f_ in F], size=16, space_after=16, spacing=1.05)
    box(s, 7.85, 1.55, 4.9, 5.15)
    label(s, 8.15, 1.78, 4.4, "Lessons & challenges", color=YELLOW)
    lessons = ["Grading free-text answers is noisy → lenient alias matching + manual spot-checks.",
               "Calibration ≠ discrimination: TS fixes the confidence scale, but barely changes which answers get trusted.",
               "Token log-probs are not offered by every API (e.g. reasoning models) → black-box signals matter.",
               "Rate limits shaped the design: resumable runs, parallel requests, caching."]
    text(s, 8.15, 2.25, 4.35, 4.35, [[("•  ", {"color": YELLOW, "bold": True}), (l_, {})] for l_ in lessons], size=14.5, space_after=14, spacing=1.05)

    # ---------------- 11. Next steps + timeline
    s = new_slide(prs, "10. Next Steps & Timeline", "We are on schedule — Week 2 milestone complete", notes=(
        "We are on schedule. Weeks 1 and 2 are done. In Weeks 3 and 4 we build Stage 2 of the pipeline: for low-agreement questions "
        "we will try a retrieval-based fact check or a multi-agent / stronger-model re-check, and we will log tokens and latency so we "
        "can compare cost against always running the expensive check. We also plan to add a harder dataset and a second model to see "
        "whether the findings generalise. Week 5 is the full evaluation against the baselines, and Week 6 is the write-up."))
    tl = [("W1", "Literature • scope", GREEN, "DONE"), ("W2", "Baselines + first results", GREEN, "DONE"),
          ("W3–4", "Build & test triage pipeline", CYAN, "NEXT"), ("W5", "Evaluate vs. baselines", MUTED, "PLANNED"),
          ("W6", "Write-up & final talk", MUTED, "PLANNED")]
    tw = 2.31
    for i, (wk, desc, c, st) in enumerate(tl):
        x = 0.6 + i * (tw + 0.15)
        box(s, x, 1.6, tw, 1.55, line=c if st != "PLANNED" else BORDER)
        text(s, x + 0.25, 1.78, tw - 0.5, 0.4, wk, size=18, bold=True, color=c if st != "PLANNED" else INK)
        text(s, x + 0.25, 2.2, tw - 0.5, 0.5, desc, size=12, color=MUTED)
        text(s, x + 0.25, 2.75, tw - 0.5, 0.3, st, size=10.5, bold=True, color=c)
    nxt = [("Build Stage 2 (escalation)", "Retrieval fact-check or multi-agent / stronger-model re-check, applied only to low-agreement questions."),
           ("Measure cost honestly", "Log tokens & latency: always-answer vs. always-escalate vs. triage; tune the agreement threshold."),
           ("Test generality", "Add a harder benchmark (e.g. SimpleQA) and a second model; larger sample for tighter estimates.")]
    nw = 3.9
    for i, (h, b) in enumerate(nxt):
        x = 0.6 + i * (nw + 0.225)
        box(s, x, 3.5, nw, 3.2)
        label(s, x + 0.3, 3.78, nw - 0.6, f"Next {i + 1}", color=CYAN)
        text(s, x + 0.3, 4.2, nw - 0.6, 0.5, h, size=18, bold=True)
        text(s, x + 0.3, 4.8, nw - 0.6, 1.8, b, size=15, color=MUTED, spacing=1.05)

    # ---------------- 12. Q&A + references
    s = new_slide(prs, None, notes="Thank you — happy to take questions.")
    text(s, 0.75, 1.0, 5.5, 0.4, "THANK YOU", size=13, color=CYAN, bold=True)
    text(s, 0.75, 1.45, 5.8, 1.0, "Questions?", size=44, bold=True)
    text(s, 0.77, 2.55, 5.4, 1.8, f"{PRESENTER} • ID {STUDENT_ID}\n{GROUP} • Seminar on AI", size=15, color=MUTED, spacing=1.15)
    box(s, 6.75, 0.9, 6.0, 5.8)
    label(s, 7.05, 1.15, 5.5, "References used this week")
    ds_ref = {"triviaqa": ("Joshi et al. (ACL 2017)", "TriviaQA: a large-scale distantly supervised challenge dataset for reading comprehension."),
              "nq_open": ("Kwiatkowski et al. (TACL 2019)", "Natural Questions: a benchmark for question answering research.")}.get(run.get("dataset"))
    refs = [("Guo et al. (ICML 2017)", "On Calibration of Modern Neural Networks — temperature scaling."),
            ("Tian et al. (EMNLP 2023)", "Just Ask for Calibration — verbalized confidence from RLHF-tuned LLMs."),
            ("Wang et al. (ICLR 2023)", "Self-Consistency Improves Chain of Thought Reasoning in Language Models."),
            ("Geifman & El-Yaniv (NeurIPS 2017)", "Selective Classification for Deep Neural Networks — risk–coverage.")]
    if ds_ref:
        refs.append(ds_ref)
    paras = []
    for a, b in refs:
        paras.append([(a, {"bold": True, "color": INK, "size": 13})])
        paras.append([(b, {"color": MUTED, "size": 12})])
    tb = text(s, 7.05, 1.65, 5.45, 4.9, paras, space_after=3)
    for j, p in enumerate(tb.text_frame.paragraphs):
        if j % 2 == 1:
            p.space_after = Pt(13)
    if tri["errors_caught_share"] is not None:
        box(s, 0.75, 4.35, 5.4, 2.35)
        label(s, 1.05, 4.6, 4.9, "Take-away", color=GREEN)
        text(s, 1.05, 5.0, 4.85, 1.6, f"{pct(tri['errors_caught_share'])} of the model's errors sit in the "
             f"{pct(tri['escalated_share'])} of questions where its own samples disagree — "
             "that is where verification should go.", size=15)

    prs.save(out_path)
    return SLIDE_COUNT[0]


def main():
    import datetime as dt

    ap = argparse.ArgumentParser()
    ap.add_argument("--results", default="results")
    ap.add_argument("--out", default="Group7_Week2_Progress_SamanRai.pptx")
    ap.add_argument("--date", default=dt.date.today().strftime("%-d %B %Y") if os.name != "nt" else dt.date.today().strftime("%#d %B %Y"))
    args = ap.parse_args()
    with open(os.path.join(args.results, "summary.json"), encoding="utf-8") as f:
        S = json.load(f)
    parts = args.date.split()
    DATE_SHORT[0] = f"{parts[0]} {parts[1][:3]} {parts[2]}" if len(parts) == 3 else args.date
    n = build(S, args.results, args.out, args.date)
    print(f"Wrote {args.out} ({n} slides). Open it, read the speaker notes, and rehearse!")


if __name__ == "__main__":
    main()
